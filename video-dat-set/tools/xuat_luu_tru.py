"""Xuất bản lưu trữ của bài giảng 3D đất sét: mỗi chương một PDF và một PPTX.

    python3 xuat_luu_tru.py [chương ...]      # mặc định cả 1–5

Ảnh slide lấy thẳng từ video đã hoàn thiện, không dựng lại: video là bản đã
kiểm, nên bản lưu trữ khớp từng điểm ảnh với cái sinh viên xem. Khung được chọn
cách cuối slide 20 khung (0,67 s): mọi khối chữ đã hiện đủ, còn 8 khung cuối
thì đang mờ dần sang slide sau (BaiGiang.tsx, biến `mo`).

PPTX mang lời giảng của từng slide trong ghi chú, lấy từ so_do.json, tức đúng
câu đã đọc (đã viết theo cách đọc, ví dụ "bốn pê" cho 4P).
"""
import io, json, math, subprocess, sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

FPS = 30
GOC = Path(__file__).resolve().parents[2]
DAY = GOC / 'video-dat-set'


def khung(video, so_khung):
    """Một khung của video, chính xác theo số thứ tự (lọc select, không tua gần đúng)."""
    ra = subprocess.run(['ffmpeg', '-loglevel', 'error', '-i', str(video), '-vf',
                         f'select=eq(n\\,{so_khung})', '-frames:v', '1', '-f', 'image2pipe',
                         '-vcodec', 'png', '-'], capture_output=True, check=True).stdout
    return Image.open(io.BytesIO(ra)).convert('RGB')


def xuat(c):
    cc = f'{c:02d}'
    dl = json.load(open(DAY / f'du-lieu/ch{cc}.json'))
    am = {(a['video'], a['so']): a for a in
          json.load(open(DAY / f'remotion/public/am-thanh/ch{cc}/so_do.json'))['slides']}
    bia = next(s for s in dl['slides'] if 'TRANG BÌA' in s['tieu_de'])
    ten = bia['tren_slide'][0]
    trang = []
    for v in (1, 2):
        sx = 2 * c - 2 + v
        video = GOC / f'chuong-{cc}/chu-de-{v}/bai-giang/EC1606_SX{sx:02d}.mp4'
        tu = 0
        for s in [s for s in dl['slides'] if s['video'] == v]:
            a = am[(v, s['so'])]
            dai = math.ceil(a['giay'] * FPS)
            anh = khung(video, tu + dai - 20)
            trang.append((anh, f"SX{sx:02d} · slide {s['so']} · {s['tieu_de']}",
                          '\n'.join(x['cau'] for x in a['cau'])))
            tu += dai
    ra = GOC / f'chuong-{cc}/slides-dat-set'
    ra.mkdir(exist_ok=True)
    goc_ten = f'EC1606_CHUONG_{c}_BAI_GIANG_DAT_SET'
    # PDF: ảnh JPEG chất lượng 90, 150 dpi cho khổ 12,8 × 7,2 in — in ra vẫn nét.
    anh = [t[0] for t in trang]
    anh[0].save(ra / f'{goc_ten}.pdf', save_all=True, append_images=anh[1:],
                quality=90, resolution=150, title=f'EC1606 · {ten}',
                author='Đỗ Thùy Hương')
    pr = Presentation()
    pr.slide_width, pr.slide_height = Emu(12192000), Emu(6858000)
    for im, nhan, loi in trang:
        sl = pr.slides.add_slide(pr.slide_layouts[6])
        b = io.BytesIO(); im.save(b, 'JPEG', quality=90); b.seek(0)
        sl.shapes.add_picture(b, 0, 0, pr.slide_width, pr.slide_height)
        sl.notes_slide.notes_text_frame.text = f'{nhan}\n\n{loi}'
    pr.core_properties.title = f'EC1606 · {ten}'
    pr.core_properties.author = 'Đỗ Thùy Hương'
    pr.save(ra / f'{goc_ten}.pptx')
    print(f'chương {c}: {len(trang)} slide → {ra.relative_to(GOC)}')


if __name__ == '__main__':
    for c in [int(x) for x in sys.argv[1:]] or range(1, 6):
        xuat(c)

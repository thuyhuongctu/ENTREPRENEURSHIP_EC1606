# -*- coding: utf-8 -*-
"""Định dạng văn bản dùng chung cho mọi tài liệu "toàn văn" trong repo này —
đúng Điều 20 Quy định đào tạo trực tuyến và xây dựng học liệu số
(QĐ 55/QĐ-ĐHSPKTVL-ĐT_ĐBCL): khổ A4, cỡ chữ 13, font Times New Roman, TCVN
6909 (Unicode — mặc định của hệ thống nên không cần xử lý riêng), giãn dòng
1,5 lines, lề trái 3cm, lề phải 2cm, lề trên 2cm, lề dưới 2,5cm.

Dùng chung cho mọi trình sinh docx/pdf trong tools/ — KHÔNG định dạng riêng
lẻ ở từng trình sinh, để đổi một chỗ là khớp toàn bộ.
"""
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_LINE_SPACING

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

PHONG_THUONG = '/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf'
PHONG_DAM = '/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf'
PHONG_NGHIENG = '/usr/share/fonts/truetype/liberation/LiberationSerif-Italic.ttf'

CO_CHU = 13
LE_TRAI, LE_PHAI, LE_TREN, LE_DUOI = 3 * cm, 2 * cm, 2 * cm, 2.5 * cm

_DA_DANG_KY_PHONG = False


def dang_ky_phong_pdf():
    global _DA_DANG_KY_PHONG
    if _DA_DANG_KY_PHONG:
        return
    pdfmetrics.registerFont(TTFont('TNR', PHONG_THUONG))
    pdfmetrics.registerFont(TTFont('TNR-Dam', PHONG_DAM))
    pdfmetrics.registerFont(TTFont('TNR-Nghieng', PHONG_NGHIENG))
    _DA_DANG_KY_PHONG = True


def tao_docx():
    """Tạo Document mới, đặt sẵn khổ A4, lề và font mặc định theo Điều 20."""
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.5)

    kieu = doc.styles['Normal']
    kieu.font.name = 'Times New Roman'
    kieu.font.size = Pt(CO_CHU)
    kieu.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return doc


def doan_van(doc, chu, dam=False, nghieng=False, co_chu=None, can_giua=False,
              cach_truoc=0, cach_sau=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(cach_truoc)
    p.paragraph_format.space_after = Pt(cach_sau)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if can_giua:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(chu)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(co_chu or CO_CHU)
    r.bold = dam
    r.italic = nghieng
    return p


class TrinhVietPdf:
    """Bọc reportlab.canvas để viết đoạn văn tự xuống dòng, đúng lề/cỡ chữ
    Điều 20, dùng chung cho mọi trình sinh PDF trong repo."""

    def __init__(self, duong_dan_ra):
        from reportlab.pdfgen import canvas
        dang_ky_phong_pdf()
        self.trang_rong, self.trang_cao = A4
        self.rong_chu = self.trang_rong - LE_TRAI - LE_PHAI
        self.c = canvas.Canvas(duong_dan_ra, pagesize=A4)
        self.y = self.trang_cao - LE_TREN
        self.duong_dan_ra = duong_dan_ra

    def _phong(self, dam, nghieng):
        if dam:
            return 'TNR-Dam'
        if nghieng:
            return 'TNR-Nghieng'
        return 'TNR'

    def _boc_dong(self, chu, ten_phong, co_chu, rong_toi_da):
        if not chu:
            return ['']
        tu = chu.split(' ')
        dong, dong_hien_tai = [], tu[0]
        for t in tu[1:]:
            thu = dong_hien_tai + ' ' + t
            if pdfmetrics.stringWidth(thu, ten_phong, co_chu) <= rong_toi_da:
                dong_hien_tai = thu
            else:
                dong.append(dong_hien_tai)
                dong_hien_tai = t
        dong.append(dong_hien_tai)
        return dong

    def viet(self, chu, dam=False, nghieng=False, co_chu=None, cach_sau=6, thut_le=0):
        co_chu = co_chu or CO_CHU
        ten_phong = self._phong(dam, nghieng)
        self.c.setFont(ten_phong, co_chu)
        giuoc_dong = co_chu * 1.5
        for dong in self._boc_dong(chu, ten_phong, co_chu, self.rong_chu - thut_le):
            if self.y < LE_DUOI:
                self.c.showPage()
                self.y = self.trang_cao - LE_TREN
                self.c.setFont(ten_phong, co_chu)
            self.c.drawString(LE_TRAI + thut_le, self.y, dong)
            self.y -= giuoc_dong
        self.y -= cach_sau

    def trang_moi(self):
        self.c.showPage()
        self.y = self.trang_cao - LE_TREN

    def luu(self):
        self.c.save()
        return self.duong_dan_ra

# -*- coding: utf-8 -*-
"""Dựng gói bài giảng từng chương theo đúng khuôn gói Chương 5 EC6035TX của
cùng tác giả (repo ComDraft, thư mục ec6035tx-chuong-5): MỘT bộ slide cho cả
chương, chia hai video (mỗi chủ đề một video), kèm kịch bản thuyết minh và
bài ôn tập tự luận hai bản (SV/GV).

Đầu vào (NGUỒN — sửa ở đây rồi chạy lại, đừng sửa tay tệp kết quả):
  docs/slide_noi_dung_chu_de.json   chữ trên slide của 10 chủ đề
  docs/kich_ban_chu_de.json         lời đọc, khớp 1-1 với slide chủ đề
  docs/goi_chuong_ch0N.json         slide khung của chương (bìa, mục tiêu,
                                    định vị, ôn tập, tình huống, tổng kết,
                                    kết) và bài ôn tập tự luận
  assets/mau/MAU_SLIDE_VLUTE.pptx   khuôn slide (xem ghi chú ở hàm main)

Đầu ra cho mỗi chương N:
  chuong-0N/slides/CHUONG_N__<TÊN>_EC1606.pptx + .pdf
  chuong-0N/kich-ban/KICH_BAN_THUYET_MINH__CHUONG_N_EC1606.docx + .pdf
  chuong-0N/on-tap/CAU_HOI_ON_TAP_TU_LUAN__CHUONG_N_EC1606 - BAN DE (SV).docx + .pdf
  chuong-0N/on-tap/CAU_HOI_ON_TAP_TU_LUAN__CHUONG_N_EC1606 - BAN CO DAP AN (GV).docx + .pdf
  chuong-0N/chu-de-M/bai-giang/NNN.jpg + video.mp4   (ảnh xuất từ chính bộ
                                    .pptx trên, video câm — xem README)

Chạy: python3 tools/gop_chu_de.py && python3 tools/dung_goi_chuong.py [1 2 ...]
Cần: python-pptx, python-docx, reportlab, Pillow, LibreOffice (soffice),
pdftoppm, ffmpeg, phông Roboto (khuôn dùng Roboto; thiếu phông thì
LibreOffice tự thay phông khác và chữ tràn khỏi khung).
"""
import copy
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Pt
from docx.shared import Pt as DocxPt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import dinh_dang_vanban as dd  # noqa: E402

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(GOC, 'docs')
KHUON = os.path.join(GOC, 'assets', 'mau', 'MAU_SLIDE_VLUTE.pptx')

MA_HP = 'EC1606'
TEN_HP = 'Khởi sự doanh nghiệp'
CHAN_TRANG = 'Học liệu số · Khởi sự doanh nghiệp (EC1606)'
GIANG_VIEN = ('Đỗ Thùy Hương – Khoa Kinh doanh và Luật, '
              'Trường Đại học Công nghệ Kỹ thuật Vĩnh Long')

# 206 từ/phút đo trên 14 bản thu thật giọng tác giả ở gói ec6035tx-chuong-5
# (ComDraft) — KHÔNG phải đo trên bộ kịch bản này. Có bản thu đầu tiên của
# EC1606 thì đo lại và sửa ở đây.
TOC_DO_DOC = 206.0
NGHI_CHUYEN_SLIDE = 1.5
THOI_LUONG_TOI_THIEU = 4.0

# Thứ tự 8 trang mẫu trong khuôn — khuôn cắt từ deck Chương 5 EC6035TX
# (giữ trang 1, 2, 3, 4, 9, 44, 49, 50 của deck gốc).
MAU_BIA, MAU_MUC_TIEU, MAU_NOI_DUNG, MAU_THE_DOC, MAU_THE_LUOI, \
    MAU_CAU_HOI, MAU_TAI_LIEU, MAU_KET = range(8)

NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
EMU = 914400

XANH_DAM = (14, 107, 45)
XANH = (74, 173, 60)
VANG = (232, 184, 0)
XAM = (150, 150, 150)
CHU = (33, 33, 33)



# ----------------------------------------------------------------------
# Nhân bản slide và thao tác chữ
# ----------------------------------------------------------------------

def nhan_ban(prs, mau):
    """Chép một trang mẫu thành trang mới ở cuối bộ — kể cả nền và ảnh.

    python-pptx không có hàm chép slide; phải chép cây hình và nối lại quan
    hệ ảnh bằng tay, nếu không ảnh trong trang mới trỏ vào rId không tồn tại
    và PowerPoint báo tệp hỏng (LibreOffice thì lặng lẽ bỏ ảnh)."""
    moi = prs.slides.add_slide(mau.slide_layout)
    for ph in list(moi.placeholders):
        ph._element.getparent().remove(ph._element)
    nen = mau._element.cSld.bg
    if nen is not None:
        moi._element.cSld.insert(0, copy.deepcopy(nen))
    doi = {}
    for rid, rel in mau.part.rels.items():
        if rel.reltype.endswith('/slideLayout') or rel.reltype.endswith('/notesSlide'):
            continue
        if rel.is_external:
            doi[rid] = moi.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            doi[rid] = moi.part.relate_to(rel.target_part, rel.reltype)
    cay = moi.shapes._spTree
    for el in mau.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'):
            continue
        cay.append(copy.deepcopy(el))
    for el in cay.iter():
        for thuoc_tinh in ('embed', 'link', 'id'):
            khoa = '{%s}%s' % (NS_R, thuoc_tinh)
            if el.get(khoa) in doi:
                el.set(khoa, doi[el.get(khoa)])
    return moi


def hinh(slide, ma):
    for sh in slide.shapes:
        if sh.shape_id == ma:
            return sh
    raise KeyError('Không có hình id=%d' % ma)


def xoa(slide, *ma):
    for m in ma:
        el = hinh(slide, m)._element
        el.getparent().remove(el)


def _dat_doan(p, chu, co=None):
    """Thay chữ của một đoạn, giữ định dạng của run đầu tiên."""
    runs = p.runs
    if not runs:
        r = p.add_run()
    else:
        r = runs[0]
        for khac in runs[1:]:
            khac._r.getparent().remove(khac._r)
    r.text = chu
    if co:
        r.font.size = Pt(co)


def dat_chu(shape, chu, co=None):
    """Đặt chữ cho hộp một đoạn (dòng tiêu đề, số trang, chân trang...).
    Nhiều dòng thì chép đoạn đầu ra thành nhiều đoạn cùng định dạng."""
    dong = chu if isinstance(chu, list) else [chu]
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    mau_p = copy.deepcopy(p0._p)
    _dat_doan(p0, dong[0], co)
    for d in dong[1:]:
        el = copy.deepcopy(mau_p)
        tf._txBody.append(el)
        _dat_doan(tf.paragraphs[-1], d, co)


def dat_the(shape, tieu_de, noi_dung, co_td=None, co_nd=None):
    """Thẻ bo góc của khuôn có hai đoạn: dòng đầu đậm màu xanh (tiêu đề),
    dòng sau là thân. Thiếu tiêu đề thì bỏ hẳn đoạn đầu, đừng để dòng trống
    — dòng trống vẫn chiếm chỗ và đẩy thân thẻ lệch xuống dưới."""
    ps = shape.text_frame.paragraphs
    if tieu_de:
        _dat_doan(ps[0], tieu_de, co_td)
    else:
        ps[0]._p.getparent().remove(ps[0]._p)
    _dat_doan(shape.text_frame.paragraphs[-1], noi_dung, co_nd)


def co_vua(chu, rong_in, cao_in, co_max, co_min, he_so_rong=0.5, gian_dong=1.25):
    """Ước cỡ chữ lớn nhất để đoạn chữ lọt khung. Không đo được chữ thật
    trong python-pptx, nên ước theo bề rộng ký tự trung bình của Roboto
    (~0,5 em chữ thường; ~0,66 em chữ in hoa đậm) — để dư một chút."""
    co = co_max
    while co > co_min:
        ky_tu_moi_dong = max(1, int(rong_in * 72 / (co * he_so_rong)))
        so_dong = 0
        for doan in (chu if isinstance(chu, list) else [chu]):
            so_dong += max(1, -(-len(doan) // ky_tu_moi_dong))
        if so_dong * co * gian_dong / 72 <= cao_in:
            return co
        co -= 1
    return co_min


def danh_so_trang(slide, so):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t == 'Học liệu số - Đào tạo từ xa':
            # Hộp chân trang của khuôn chỉ rộng 5,44 inch — vừa câu gốc
            # "Học liệu số - Đào tạo từ xa" nhưng cắt mất "(EC1606)" của
            # câu mới; nới tới sát ô số trang (x = 9,67 inch).
            dat_chu(sh, CHAN_TRANG)
            sh.width = int(8.8 * EMU)
        elif t.isdigit() and sh.top > 9.5 * EMU:
            dat_chu(sh, str(so))


def ghi_loi_doc(slide, loi_doc):
    slide.notes_slide.notes_text_frame.text = loi_doc


def hoa(chu):
    return chu.upper()


# ----------------------------------------------------------------------
# Khung minh họa bên phải (thay ảnh trong khung bo góc của trang mẫu)
# ----------------------------------------------------------------------

def _phong(kieu, co):
    """Tìm tệp Roboto qua fontconfig thay vì ghi cứng đường dẫn — mỗi máy
    cài Roboto một chỗ, tên tệp nghiêng cũng khác nhau (Roboto-Italic.ttf
    hay Roboto-RegularItalic.ttf). Không thấy Roboto thì dừng hẳn: fc-match
    lặng lẽ trả phông khác, ảnh ra vẫn "đúng" nhưng lệch hẳn kiểu chữ khuôn."""
    tep = subprocess.run(['fc-match', '-f', '%{file}', 'Roboto:style=' + kieu],
                         capture_output=True, text=True).stdout
    assert 'Roboto' in os.path.basename(tep), 'Thiếu phông Roboto (%s) — xem README' % kieu
    return ImageFont.truetype(tep, co)


def _boc(draw, chu, phong, rong):
    tu, dong, hien = chu.split(), [], ''
    for t in tu:
        thu = (hien + ' ' + t).strip()
        if draw.textlength(thu, font=phong) <= rong:
            hien = thu
        else:
            dong.append(hien)
            hien = t
    if hien:
        dong.append(hien)
    return dong


def ve_khung_dinh_vi(tieu_de, muc, chi_so_hien_tai, chu_thich):
    """Khung "đang ở đâu" bên phải mỗi slide: một trục dọc liệt kê các mục
    của chủ đề (hoặc của chương), mục đã qua tick xanh, mục đang học chấm
    vàng, mục sắp học xám — cùng kiểu khung "Chúng ta đang ở đâu" của gói
    Chương 5 EC6035TX. Trả về PNG dạng bytes.

    Khung trong khuôn có tỉ lệ 4481564 x 5827243 EMU; ảnh vẽ đúng tỉ lệ ấy,
    lệch tỉ lệ thì LibreOffice kéo giãn chữ trong ảnh."""
    W, H = 896, 1165
    im = Image.new('RGB', (W, H), 'white')
    d = ImageDraw.Draw(im)
    # Khung chỉ chiếm ~4,5 inch trên slide 20 inch: chữ dưới ~36 px ở đây
    # thành chữ 7-8 pt khi chiếu, không đọc nổi (lần dựng đầu đã như thế).
    p_td = _phong('Bold', 46)
    p_muc = _phong('Medium', 40)
    p_muc_dam = _phong('Bold', 41)
    p_ct = _phong('Italic', 30)
    p_nhan = _phong('Regular', 28)
    y = 46
    for dong in _boc(d, tieu_de, p_td, W - 100):
        w = d.textlength(dong, font=p_td)
        d.text(((W - w) / 2, y), dong, font=p_td, fill=XANH_DAM)
        y += 56
    y += 10
    d.line((60, y, W - 60, y), fill=XANH, width=3)
    y += 34

    # Quá nhiều mục thì chỉ hiện một cửa sổ quanh mục hiện tại, kèm dấu "…"
    # ở đầu/cuối — liệt kê đủ 15 mục thì chữ phải nhỏ tới mức không đọc nổi.
    TOI_DA = 6
    dau = 0
    if len(muc) > TOI_DA:
        dau = min(max(0, chi_so_hien_tai - TOI_DA // 2), len(muc) - TOI_DA)
    hien = list(range(dau, min(len(muc), dau + TOI_DA)))
    x_truc = 92
    vung_cao = H - 120 - y
    buoc = vung_cao / max(len(hien), 1)
    toa_do = []
    for k, i in enumerate(hien):
        toa_do.append((i, y + k * buoc + 22))
    if len(toa_do) > 1:
        d.line((x_truc, toa_do[0][1], x_truc, toa_do[-1][1]), fill=(200, 222, 196), width=5)
    if dau > 0:
        d.text((x_truc - 10, y - 30), '…', font=p_muc, fill=XAM)
    for i, yy in toa_do:
        r = 20
        if i < chi_so_hien_tai:
            d.ellipse((x_truc - r, yy - r, x_truc + r, yy + r), fill=XANH)
            d.line((x_truc - 9, yy, x_truc - 2, yy + 8, x_truc + 10, yy - 8), fill='white', width=5)
            mau, phong, nhan = XANH_DAM, p_muc, 'Đã học'
        elif i == chi_so_hien_tai:
            d.ellipse((x_truc - r - 4, yy - r - 4, x_truc + r + 4, yy + r + 4), fill=VANG)
            mau, phong, nhan = CHU, p_muc_dam, 'Đang học'
        else:
            d.ellipse((x_truc - r, yy - r, x_truc + r, yy + r), outline=XAM, width=4, fill='white')
            mau, phong, nhan = XAM, p_muc, 'Sắp học'
        dong = _boc(d, muc[i], phong, W - x_truc - 80)[:2]
        if len(_boc(d, muc[i], phong, W - x_truc - 80)) > 2:
            dong[1] = dong[1].rstrip('.,;:') + '…'
        yt = yy - 24
        for dg in dong:
            d.text((x_truc + 44, yt), dg, font=phong, fill=mau)
            yt += 48
        d.text((x_truc + 44, yt + 2), nhan, font=p_nhan, fill=XAM if mau == XAM else XANH)
    if hien and hien[-1] < len(muc) - 1:
        d.text((x_truc - 10, toa_do[-1][1] + 40), '…', font=p_muc, fill=XAM)
    for k, dong in enumerate(_boc(d, chu_thich, p_ct, W - 100)[:2]):
        w = d.textlength(dong, font=p_ct)
        d.text(((W - w) / 2, H - 100 + k * 38), dong, font=p_ct, fill=XAM)
    buf = io.BytesIO()
    im.save(buf, 'PNG', optimize=True)
    buf.seek(0)
    return buf


def thay_anh_khung(slide, ma_khung, anh):
    _, rid = slide.part.get_or_add_image_part(anh)
    blip = hinh(slide, ma_khung)._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
    blip.set('{%s}embed' % NS_R, rid)


# ----------------------------------------------------------------------
# Các kiểu trang
# ----------------------------------------------------------------------

def tach_y(gach):
    """Tách một gạch đầu dòng "Nhãn: nội dung" thành (tiêu đề thẻ, thân thẻ).
    Chỉ tách khi phần nhãn ngắn — tách câu dài ở dấu hai chấm giữa câu thì
    tiêu đề thẻ thành nửa câu cụt."""
    for dau in (': ', ' — ', ' – '):
        if dau in gach:
            trai, phai = gach.split(dau, 1)
            if 1 <= len(trai.split()) <= 7 and len(phai.split()) >= 3:
                return trai.strip(), phai[0].upper() + phai[1:]
    return '', gach


def trang_the(prs, khuon, tieu_de, the, khung_anh, loi_doc):
    """Trang thẻ nội dung: dải tiêu đề xanh, các thẻ bên trái, khung minh
    họa bên phải. 4 thẻ dùng lưới 2x2 (trang mẫu 9), còn lại xếp hàng dọc
    (trang mẫu 4) — giãn chiều cao thẻ theo số thẻ."""
    n = len(the)
    mau = khuon.slides[MAU_THE_LUOI if n == 4 else MAU_THE_DOC]
    s = nhan_ban(prs, mau)
    co_td = min(44, int(13.0 * 72 / (0.68 * max(len(tieu_de), 1))))
    dat_chu(hinh(s, 218), hoa(tieu_de), co=max(co_td, 22))
    if n == 4:
        o = [hinh(s, m) for m in (237, 238, 239, 240)]
        for sh, (td, nd) in zip(o, the):
            ct = co_vua(nd, 5.9, 2.9 - (0.55 if td else 0), 22, 15)
            dat_the(sh, td, nd, co_td=min(24, ct + 2), co_nd=ct)
    else:
        mau_the = hinh(s, 237)._element
        for m in (237, 238, 239):
            xoa(s, m)
        tren, duoi, khe = 3.45, 10.05, 0.18
        cao = (duoi - tren - khe * (n - 1)) / n
        for k, (td, nd) in enumerate(the):
            el = copy.deepcopy(mau_the)
            s.shapes._spTree.append(el)
            sh = s.shapes[-1]
            sh.top = int((tren + k * (cao + khe)) * EMU)
            sh.height = int(cao * EMU)
            ct = co_vua(nd, 12.6, cao - 0.3 - (0.45 if td else 0), 22, 14)
            dat_the(sh, td, nd, co_td=min(24, ct + 2), co_nd=ct)
    thay_anh_khung(s, 233, khung_anh)
    ghi_loi_doc(s, loi_doc)
    return s


def trang_bia(prs, khuon, so, ten_chuong, loi_doc):
    s = nhan_ban(prs, khuon.slides[MAU_BIA])
    td = 'CHƯƠNG %d: %s' % (so, hoa(ten_chuong))
    dat_chu(hinh(s, 109), td, co=min(40, int(15.5 * 72 / (0.66 * len(td)) * 2)))
    dat_chu(hinh(s, 112), 'Học phần: %s (%s)' % (TEN_HP, MA_HP))
    # Thay hình laptop + mũi tên của khuôn bằng Hương AI đứng giảng.
    xoa(s, 105, 85, 88)
    them_huong_ai(s, ANH_HUONG_AI_GIANG, 13.5, 5.9, rong=5.3, khung=True)
    ghi_loi_doc(s, loi_doc)
    return s


ANH_HUONG_AI_GIANG = os.path.join(GOC, 'assets', 'huong-ai', 'HUONG_AI_GIANG_BAI.jpg')
ANH_HUONG_AI_DUNG = os.path.join(GOC, 'assets', 'huong-ai', 'HUONG_AI_DUNG.jpg')
CHU_THICH_HUONG_AI = 'Hương AI – nhân vật minh họa do AI tạo, không phải ảnh thật'


def them_huong_ai(slide, anh, x, y, rong=None, cao=None, khung=False):
    """Đặt hình nhân vật "Hương AI" kèm chú thích ngay dưới.

    Chú thích là bắt buộc, đừng bỏ: ở ComDraft, ảnh nhân vật do AI dựng
    (áo dài) từng bị gỡ khỏi ba bài thực hành theo yêu cầu của cô Hương vì
    dễ bị hiểu nhầm là ảnh thật của giảng viên."""
    from pptx.dml.color import RGBColor
    w, h = Image.open(anh).size
    if rong:
        cao = rong * h / w
    else:
        rong = cao * w / h
    pic = slide.shapes.add_picture(anh, int(x * EMU), int(y * EMU), int(rong * EMU), int(cao * EMU))
    if khung:
        pic.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        pic.line.width = Pt(3)
    hop = slide.shapes.add_textbox(int((x - 1.0) * EMU), int((y + cao + 0.02) * EMU),
                                   int((rong + 2.0) * EMU), int(0.3 * EMU))
    p = hop.text_frame.paragraphs[0]
    p.alignment = 2  # giữa
    r = p.add_run()
    r.text = CHU_THICH_HUONG_AI
    r.font.size = Pt(11)
    r.font.italic = True
    r.font.name = 'Roboto'
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if khung else RGBColor(0x75, 0x75, 0x75)
    return pic


def trang_muc_tieu(prs, khuon, so, muc_tieu, loi_doc):
    s = nhan_ban(prs, khuon.slides[MAU_MUC_TIEU])
    dat_chu(hinh(s, 141), 'MỤC TIÊU CHƯƠNG %d' % so)
    for ma, chu in zip((142, 143, 144, 146), muc_tieu):
        hinh(s, ma).width = int(7.0 * EMU)
        dat_chu(hinh(s, ma), chu, co=co_vua(chu, 7.0, 0.95, 22, 15))
    # Hương AI đứng thay hình bia bắn tên của khuôn (thẳng hàng với cột phải).
    xoa(s, 147)
    them_huong_ai(s, ANH_HUONG_AI_DUNG, 12.9, 3.75, cao=6.0, khung=False)
    ghi_loi_doc(s, loi_doc)
    return s


def trang_noi_dung(prs, khuon, so, muc, loi_doc):
    """Trang mẫu có 5 mục đánh số; chương EC1606 chỉ có 3 (hai chủ đề và
    phần ôn tập) nên xóa hai mục cuối thay vì để ô trống."""
    s = nhan_ban(prs, khuon.slides[MAU_NOI_DUNG])
    dat_chu(hinh(s, 191), 'NỘI DUNG CHƯƠNG %d' % so)
    for ma, chu in zip((193, 198, 199), muc):
        dat_chu(hinh(s, ma), chu, co=co_vua(chu, 7.6, 1.05, 22, 15))
        hinh(s, ma).width = int(7.8 * EMU)
    xoa(s, 182, 185, 196, 197, 200, 201)
    ghi_loi_doc(s, loi_doc)
    return s


def trang_cau_hoi(prs, khuon, so, cau, khung_anh, loi_doc):
    s = nhan_ban(prs, khuon.slides[MAU_CAU_HOI])
    dat_chu(hinh(s, 218), 'CÂU HỎI ÔN TẬP CHƯƠNG %d' % so)
    for ma, chu in zip((238, 240, 242, 244, 246), cau):
        dat_chu(hinh(s, ma), chu, co=co_vua(chu, 12.2, 1.0, 22, 15))
    thay_anh_khung(s, 233, khung_anh)
    ghi_loi_doc(s, loi_doc)
    return s


def trang_tai_lieu(prs, khuon, loi_doc):
    """Đúng mục 10 đề cương EC1606. Đề cương ghi [2] trùng hệt [1] (cùng
    giáo trình) — giữ đúng số thứ tự đề cương, chỉ không chép lại [2]."""
    s = nhan_ban(prs, khuon.slides[MAU_TAI_LIEU])
    dat_chu(hinh(s, 304), '[1] Nguyễn Thị Ngọc Huyền và Ngô Thị Việt Nga (2024), '
            'Giáo trình khởi sự kinh doanh (tái bản lần thứ hai), NXB Đại học Kinh tế Quốc dân.')
    dat_chu(hinh(s, 305), '[3] Trường Đại học Bách khoa Hà Nội (2023), Tài liệu tư vấn: '
            'Kỹ năng kinh doanh cho phụ nữ và người dân tộc thiểu số.')
    ghi_loi_doc(s, loi_doc)
    return s


def trang_ket(prs, khuon, loi_doc):
    s = nhan_ban(prs, khuon.slides[MAU_KET])
    dat_chu(hinh(s, 482), 'CẢM ƠN CÁC EM ĐÃ THEO DÕI')
    # Câu gốc của khuôn ("HỌC LIỆU SỐ - ĐÀO TẠO TỪ XA") ngắn hơn; giữ cỡ chữ
    # gốc thì câu này xuống hai dòng và đè lên dòng tên khoa ngay dưới.
    dat_chu(hinh(s, 483), 'HỌC LIỆU SỐ · KHỞI SỰ DOANH NGHIỆP', co=30)
    ghi_loi_doc(s, loi_doc)
    return s


# ----------------------------------------------------------------------
# Dựng một chương
# ----------------------------------------------------------------------

def khong_dau(chu):
    chu = chu.replace('đ', 'd').replace('Đ', 'D')
    chu = unicodedata.normalize('NFD', chu)
    chu = ''.join(c for c in chu if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^A-Za-z0-9]+', '_', chu).strip('_').upper()


def thoi_luong(loi_doc):
    return max(THOI_LUONG_TOI_THIEU, len(loi_doc.split()) / TOC_DO_DOC * 60.0 + NGHI_CHUYEN_SLIDE)


def dung_chuong(so, slides_cd, kich_ban, goi):
    ma = 'ch%02d' % so
    thu_muc = os.path.join(GOC, 'chuong-%02d' % so)
    # Trang mẫu phải lấy từ CHÍNH bộ đang dựng: chép từ một bản khuôn nạp
    # riêng thì layout/master của bản kia bị kéo theo, tệp ra mang hai bộ
    # master trùng tên và LibreOffice báo "source file could not be loaded".
    prs = Presentation(KHUON)
    khuon = prs
    so_mau = len(prs.slides)
    them = goi['slide_them']
    cd = [kich_ban['%s_cd%d' % (ma, m)] for m in (1, 2)]
    sl = [slides_cd['%s_cd%d' % (ma, m)] for m in (1, 2)]
    for m in (0, 1):
        assert len(cd[m]['diem_doc']) == len(sl[m]), '%s_cd%d: lời đọc và slide lệch số' % (ma, m + 1)
    ten_cd = ['%d.%d. %s' % (so, m + 1, cd[m]['tieu_de']) for m in (0, 1)]
    muc_chuong = [ten_cd[0], ten_cd[1], 'Câu hỏi ôn tập', 'Tình huống minh họa', 'Tổng kết chương']

    # (tiêu đề hiển thị, lời đọc, video số mấy, chữ trên slide, gợi ý hình)
    # — hai phần cuối chỉ dùng cho tệp "nội dung slide" (xem viet_noi_dung_slide).
    ban_ghi = []

    def ghi(tieu_de, loi_doc, video, noi_dung=(), goi_y_hinh=''):
        ban_ghi.append((tieu_de, loi_doc, video, list(noi_dung), goi_y_hinh))

    def the_ra_dong(the):
        return [('%s: %s' % (td, nd)) if td else nd for td, nd in the]

    trang_bia(prs, khuon, so, goi['ten_chuong'], them['bia']['loi_doc'])
    ghi('TRANG BÌA', them['bia']['loi_doc'], 1,
        ['CHƯƠNG %d: %s' % (so, hoa(goi['ten_chuong'])), 'Học phần: %s (%s)' % (TEN_HP, MA_HP),
         'Khoa: Kinh doanh và Luật', 'Trường Đại học Công nghệ Kỹ thuật Vĩnh Long'],
        'Hương AI đứng giảng (assets/huong-ai/HUONG_AI_GIANG_BAI.jpg) + chú thích "nhân vật minh họa do AI tạo"')
    trang_muc_tieu(prs, khuon, so, goi['muc_tieu'], them['muc_tieu']['loi_doc'])
    ghi('MỤC TIÊU CHƯƠNG %d' % so, them['muc_tieu']['loi_doc'], 1, goi['muc_tieu'],
        'Hương AI đứng toàn thân (assets/huong-ai/HUONG_AI_DUNG.jpg) + chú thích; dấu tick trước mỗi mục tiêu')
    trang_noi_dung(prs, khuon, so, [ten_cd[0], ten_cd[1], 'Ôn tập – Tình huống – Tổng kết chương'],
                   them['noi_dung']['loi_doc'])
    ghi('NỘI DUNG CHƯƠNG %d' % so, them['noi_dung']['loi_doc'], 1,
        ['1. ' + ten_cd[0], '2. ' + ten_cd[1], '3. Ôn tập – Tình huống – Tổng kết chương'],
        'Danh sách đánh số 1-2-3')

    for m in (0, 1):
        video = m + 1
        if m == 1:
            dv = them['dang_o_dau']
            anh = ve_khung_dinh_vi('Chương %d — chúng ta đang ở đâu' % so, muc_chuong, 1,
                                   'Video 2 bắt đầu từ chủ đề %d.2' % so)
            trang_the(prs, khuon, 'Chúng ta đang ở đâu',
                      [(c['tieu_de'], c['noi_dung']) for c in dv['the']], anh, dv['loi_doc'])
            ghi('CHÚNG TA ĐANG Ở ĐÂU', dv['loi_doc'], 2,
                the_ra_dong([(c['tieu_de'], c['noi_dung']) for c in dv['the']]),
                'Trục tiến độ chương: chủ đề %d.1 đã học, %d.2 đang học, ôn tập sắp học' % (so, so))
        tieu_de_muc = [x['title'] for x in sl[m]]
        for i, (x, dd_) in enumerate(zip(sl[m], cd[m]['diem_doc'])):
            td = ten_cd[m] if i == 0 else x['title']
            the = [tach_y(g) for g in x['bullets']]
            anh = ve_khung_dinh_vi(x['kicker'] or ('Chủ đề %d.%d' % (so, m + 1)), tieu_de_muc, i,
                                   'Chủ đề %d.%d · slide %d/%d' % (so, m + 1, i + 1, len(sl[m])))
            trang_the(prs, khuon, td, the, anh, dd_['loi_doc'])
            ghi(hoa(td), dd_['loi_doc'], video, the_ra_dong(the),
                'Trục "đang ở đâu" của chủ đề %d.%d (mục %d/%d); nhãn nhỏ: %s' % (
                    so, m + 1, i + 1, len(sl[m]), x['kicker'] or '—'))

    ch = them['cau_hoi_on_tap']
    trang_cau_hoi(prs, khuon, so, ch['cau'],
                  ve_khung_dinh_vi('Nhớ lại toàn chương', muc_chuong, 2, 'Tự trả lời trước khi xem lại video'),
                  ch['loi_doc'])
    ghi('CÂU HỎI ÔN TẬP CHƯƠNG %d' % so, ch['loi_doc'], 2,
        ['%d. %s' % (k, c) for k, c in enumerate(ch['cau'], 1)], 'Năm câu đánh số')
    th = them['tinh_huong']
    trang_the(prs, khuon, 'Tình huống minh họa', [(c['tieu_de'], c['noi_dung']) for c in th['the']],
              ve_khung_dinh_vi('Vận dụng vào tình huống', muc_chuong, 3,
                               'Ví dụ dựng ra để giảng dạy, không phải số liệu thật'),
              th['loi_doc'])
    ghi('TÌNH HUỐNG MINH HỌA', th['loi_doc'], 2,
        the_ra_dong([(c['tieu_de'], c['noi_dung']) for c in th['the']]),
        'Bốn thẻ; ghi rõ "ví dụ dựng ra để giảng dạy, không phải số liệu thật"')
    tk = them['tong_ket']
    trang_the(prs, khuon, 'Tổng kết chương %d' % so, [(c['tieu_de'], c['noi_dung']) for c in tk['the']],
              ve_khung_dinh_vi('Chương %d — đã đi hết' % so, muc_chuong, 4, 'Làm bài tập và bài ôn tập tự luận'),
              tk['loi_doc'])
    ghi('TỔNG KẾT CHƯƠNG %d' % so, tk['loi_doc'], 2,
        the_ra_dong([(c['tieu_de'], c['noi_dung']) for c in tk['the']]), 'Bốn thẻ')
    trang_tai_lieu(prs, khuon, them['tai_lieu']['loi_doc'])
    ghi('TÀI LIỆU HỌC TẬP', them['tai_lieu']['loi_doc'], 2, [
        'Giáo trình chính: [1] Nguyễn Thị Ngọc Huyền và Ngô Thị Việt Nga (2024), Giáo trình khởi sự kinh '
        'doanh (tái bản lần thứ hai), NXB Đại học Kinh tế Quốc dân.',
        'Tài liệu tham khảo: [3] Trường Đại học Bách khoa Hà Nội (2023), Tài liệu tư vấn: Kỹ năng kinh '
        'doanh cho phụ nữ và người dân tộc thiểu số.'], 'Hai khối: giáo trình chính | tài liệu tham khảo')
    trang_ket(prs, khuon, them['ket']['loi_doc'])
    ghi('KẾT THÚC CHƯƠNG', them['ket']['loi_doc'], 2,
        ['CẢM ƠN CÁC EM ĐÃ THEO DÕI', 'HỌC LIỆU SỐ · KHỞI SỰ DOANH NGHIỆP',
         'Khoa Kinh doanh và Luật – Trường Đại học Công nghệ Kỹ thuật Vĩnh Long',
         'fel@vlute.edu.vn · https://vlute.edu.vn'], 'Nền xanh đậm như trang bìa')

    # Bỏ 8 trang mẫu ở đầu (đã chép xong) — python-pptx không lưu phần ảnh
    # nào không còn ai trỏ tới, nên tệp ra không mang theo ảnh thừa.
    lst = prs.slides._sldIdLst
    for el in list(lst)[:so_mau]:
        prs.part.drop_rel(el.get('{%s}id' % NS_R))
        lst.remove(el)
    for k, s in enumerate(prs.slides, 1):
        danh_so_trang(s, k)

    viet_noi_dung_slide(so, goi, ban_ghi, thu_muc)
    if CHI_VAN_BAN:
        viet_kich_ban(so, goi, ban_ghi, thu_muc)
        viet_on_tap(so, goi, thu_muc)
        return ban_ghi

    ten_tep = 'CHUONG_%d__%s_%s' % (so, khong_dau(goi['ten_chuong']), MA_HP)
    tm_slide = os.path.join(thu_muc, 'slides')
    os.makedirs(tm_slide, exist_ok=True)
    for cu in os.listdir(tm_slide):
        os.remove(os.path.join(tm_slide, cu))
    duong_pptx = os.path.join(tm_slide, ten_tep + '.pptx')
    prs.save(duong_pptx)
    duong_pdf = xuat_pdf(duong_pptx, tm_slide)
    xuat_anh_va_video(duong_pdf, thu_muc, ban_ghi)
    viet_kich_ban(so, goi, ban_ghi, thu_muc)
    viet_on_tap(so, goi, thu_muc)
    return ban_ghi


def xuat_pdf(duong_pptx, thu_muc_ra):
    """LibreOffice từng có lần trả mã 0 mà không ra tệp (xem CLAUDE.md của
    ComDraft) — nên kiểm có tệp PDF thật, đừng tin mã thoát."""
    with tempfile.TemporaryDirectory() as tam:
        subprocess.run(['soffice', '-env:UserInstallation=file://%s/hs' % tam, '--headless',
                        '--convert-to', 'pdf', '--outdir', tam, duong_pptx],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        ra = os.path.join(tam, os.path.basename(duong_pptx)[:-5] + '.pdf')
        assert os.path.exists(ra), 'soffice không xuất được PDF cho %s' % duong_pptx
        dich = os.path.join(thu_muc_ra, os.path.basename(ra))
        shutil.move(ra, dich)
    return dich


def xuat_anh_va_video(duong_pdf, thu_muc, ban_ghi):
    with tempfile.TemporaryDirectory() as tam:
        # 96 dpi trên trang 20 x 11,25 inch = đúng 1920 x 1080.
        subprocess.run(['pdftoppm', '-r', '96', '-jpeg', '-jpegopt', 'quality=90',
                        duong_pdf, os.path.join(tam, 's')], check=True)
        anh = sorted(f for f in os.listdir(tam) if f.endswith('.jpg'))
        assert len(anh) == len(ban_ghi), 'PDF có %d trang, deck có %d slide' % (len(anh), len(ban_ghi))
        for video in (1, 2):
            tm = os.path.join(thu_muc, 'chu-de-%d' % video, 'bai-giang')
            os.makedirs(tm, exist_ok=True)
            for cu in os.listdir(tm):
                os.remove(os.path.join(tm, cu))
            chi_so = [i for i, b in enumerate(ban_ghi) if b[2] == video]
            danh_sach = os.path.join(tam, 'v%d.txt' % video)
            with open(danh_sach, 'w') as f:
                for k, i in enumerate(chi_so, 1):
                    dich = os.path.join(tm, '%03d.jpg' % k)
                    shutil.copy(os.path.join(tam, anh[i]), dich)
                    f.write("file '%s'\nduration %.2f\n" % (dich, thoi_luong(ban_ghi[i][1])))
                # concat demuxer bỏ qua "duration" của dòng cuối nếu không
                # nhắc lại tệp cuối — thiếu dòng này slide kết chỉ hiện 1 khung.
                f.write("file '%s'\n" % dich)
            subprocess.run(['ffmpeg', '-y', '-loglevel', 'error', '-f', 'concat', '-safe', '0',
                            '-i', danh_sach, '-vf', 'fps=25,format=yuv420p', '-c:v', 'libx264',
                            '-preset', 'veryfast', '-crf', '26', '-movflags', '+faststart',
                            os.path.join(tm, 'video.mp4')], check=True)


# ----------------------------------------------------------------------
# Văn bản đi kèm: kịch bản thuyết minh, bài ôn tập tự luận
# ----------------------------------------------------------------------

class HaiBan:
    """Viết song song docx và pdf cùng một nội dung, đúng định dạng Điều 20
    (dinh_dang_vanban.py) — cùng cách các trình sinh khác trong tools/."""

    def __init__(self, duong_khong_duoi):
        self.doc = dd.tao_docx()
        self.pdf = dd.TrinhVietPdf(duong_khong_duoi + '.pdf')
        self.duong = duong_khong_duoi

    def doan(self, chu, dam=False, nghieng=False, co=None, giua=False, thut=0, sau=6):
        p = dd.doan_van(self.doc, chu, dam=dam, nghieng=nghieng, co_chu=co, can_giua=giua, cach_sau=sau)
        if thut:
            p.paragraph_format.left_indent = DocxPt(thut)
        self.pdf.viet(chu, dam=dam, nghieng=nghieng, co_chu=co, cach_sau=sau, thut_le=thut)

    def bang(self, tieu_de_cot, hang):
        t = self.doc.add_table(rows=1, cols=len(tieu_de_cot))
        t.style = 'Table Grid'
        for o, chu in zip(t.rows[0].cells, tieu_de_cot):
            o.text = chu
            o.paragraphs[0].runs[0].bold = True
        for h in hang:
            for o, chu in zip(t.add_row().cells, h):
                o.text = chu
        for h in hang:
            self.pdf.viet(' | '.join(h), cach_sau=2)
        self.pdf.viet('', cach_sau=4)

    def luu(self):
        self.doc.save(self.duong + '.docx')
        self.pdf.luu()


def doc_thoi_luong(giay):
    giay = int(round(giay / 5.0) * 5)
    if giay < 60:
        return '~%d giây' % giay
    return '~%d phút %02d' % (giay // 60, giay % 60) if giay % 60 else '~%d phút' % (giay // 60)


def viet_noi_dung_slide(so, goi, ban_ghi, thu_muc):
    """Tệp chữ thuần, không thiết kế: nội dung từng slide (tiêu đề, các ý,
    gợi ý hình, lời đọc) theo đúng thứ tự bộ slide — để tác giả tự dựng
    slide bằng công cụ khác (Canva, PowerPoint) mà vẫn khớp kịch bản và
    video. Cô Hương yêu cầu tệp này để tự làm slide."""
    tm = os.path.join(thu_muc, 'noi-dung-slide')
    os.makedirs(tm, exist_ok=True)
    v = HaiBan(os.path.join(tm, 'NOI_DUNG_SLIDE__CHUONG_%d_%s' % (so, MA_HP)))
    v.doan('NỘI DUNG SLIDE BÀI GIẢNG', dam=True, co=15, giua=True)
    v.doan('CHƯƠNG %d: %s' % (so, hoa(goi['ten_chuong'])), dam=True, giua=True)
    v.doan('Học phần: %s (%s) – Học liệu số, giảng dạy trực tuyến' % (TEN_HP, MA_HP), giua=True)
    v.doan('Giảng viên biên soạn: ' + GIANG_VIEN, nghieng=True, giua=True)
    v.doan('CÁCH DÙNG', dam=True, sau=4)
    for y in ('Mỗi mục dưới đây là một slide, đúng thứ tự trình chiếu; %d slide, chia hai video.' % len(ban_ghi),
              '"Trên slide": chữ đưa lên slide (mỗi gạch là một ý/thẻ; phần trước dấu hai chấm là tiêu đề thẻ).',
              '"Gợi ý hình": cách trình bày dùng ở bộ slide dựng sẵn — tự do thay đổi.',
              '"Lời đọc": lời thuyết minh khớp với slide — trùng với kịch bản thuyết minh của chương.',
              'Quy định (Điều 18): tối đa 10 dòng/slide, 20 từ/dòng; có tên trường, khoa trên mỗi slide; '
              'chân trang "Học liệu số · Khởi sự doanh nghiệp (EC1606)".'):
        v.doan('• ' + y, thut=12, sau=2)
    video_truoc = 0
    for k, (td, loi, video, noi_dung, goi_y) in enumerate(ban_ghi, 1):
        if video != video_truoc:
            v.doan('VIDEO %d' % video, dam=True, co=14, sau=4)
            video_truoc = video
        v.doan('SLIDE %d – %s' % (k, td), dam=True, sau=2)
        v.doan('Trên slide:', nghieng=True, sau=1)
        for y in noi_dung:
            v.doan('– ' + y, thut=12, sau=1)
        if goi_y:
            v.doan('Gợi ý hình: ' + goi_y, nghieng=True, sau=2)
        v.doan('Lời đọc (%s): %s' % (doc_thoi_luong(thoi_luong(loi)), loi), co=12, sau=10)
    v.luu()


def viet_kich_ban(so, goi, ban_ghi, thu_muc):
    tm = os.path.join(thu_muc, 'kich-ban')
    os.makedirs(tm, exist_ok=True)
    v = HaiBan(os.path.join(tm, 'KICH_BAN_THUYET_MINH__CHUONG_%d_%s' % (so, MA_HP)))
    v.doan('KỊCH BẢN THUYẾT MINH VIDEO BÀI GIẢNG', dam=True, co=15, giua=True)
    v.doan('CHƯƠNG %d: %s' % (so, hoa(goi['ten_chuong'])), dam=True, giua=True)
    v.doan('Học phần: %s (%s) – Học liệu số, giảng dạy trực tuyến' % (TEN_HP, MA_HP), giua=True)
    v.doan('Giảng viên thuyết minh: ' + GIANG_VIEN, nghieng=True, giua=True)
    v.doan('Kịch bản cho bộ slide %d trang, chia hai video theo hai chủ đề của chương. '
           'Lời đọc của từng slide cũng nằm trong phần ghi chú (notes) của tệp .pptx.' % len(ban_ghi))
    v.doan('HƯỚNG DẪN SỬ DỤNG KỊCH BẢN', dam=True, sau=4)
    for y in ('Nghỉ ngắn khoảng nửa giây sau mỗi ý; nghỉ một giây trước khi chuyển slide.',
              'Số trong lời đọc đã viết thành chữ để đọc trôi; trên slide vẫn để dạng số.',
              'Thời lượng ghi ở mỗi slide tính từ số từ thật của lời đọc, theo tốc độ %d từ mỗi phút '
              '— tốc độ đo trên bản thu thật của cùng tác giả ở học phần EC6035TX, CHƯA đo trên bản thu '
              'của học phần này. Thu xong video đầu tiên nên đối chiếu lại.' % TOC_DO_DOC,
              'Các ví dụ "chị Hoa" và "gốm đỏ Vĩnh Long" là tình huống dựng ra để giảng dạy, không phải '
              'số liệu thật — khi đọc giữ nguyên câu nói rõ điều này.'):
        v.doan('• ' + y, thut=12, sau=2)
    v.doan('PHÂN CHIA VIDEO', dam=True, sau=4)
    hang = []
    for video in (1, 2):
        chi_so = [i for i, b in enumerate(ban_ghi) if b[2] == video]
        tong = sum(thoi_luong(ban_ghi[i][1]) for i in chi_so)
        tu = sum(len(ban_ghi[i][1].split()) for i in chi_so)
        noi_dung = ('Mở chương • Chủ đề %d.1' % so) if video == 1 else \
            ('Định vị • Chủ đề %d.2 • Ôn tập, tình huống, tổng kết chương' % so)
        hang.append(('VIDEO %d' % video, noi_dung, '%d – %d' % (chi_so[0] + 1, chi_so[-1] + 1),
                     '%s (%d từ)' % (doc_thoi_luong(tong), tu)))
    v.bang(('Video', 'Nội dung', 'Slide', 'Thời lượng ước tính'), hang)
    video_truoc = 0
    for k, (td, loi, video, _, _) in enumerate(ban_ghi, 1):
        if video != video_truoc:
            v.doan('VIDEO %d' % video, dam=True, co=14, sau=4)
            video_truoc = video
        v.doan('SLIDE %d – %s   (%s)' % (k, td, doc_thoi_luong(thoi_luong(loi))), dam=True, sau=2)
        v.doan(loi, sau=8)
    v.luu()


def viet_on_tap(so, goi, thu_muc):
    tm = os.path.join(thu_muc, 'on-tap')
    os.makedirs(tm, exist_ok=True)
    goc = os.path.join(tm, 'CAU_HOI_ON_TAP_TU_LUAN__CHUONG_%d_%s' % (so, MA_HP))
    for ban_gv in (False, True):
        v = HaiBan(goc + (' - BAN CO DAP AN (GV)' if ban_gv else ' - BAN DE (SV)'))
        v.doan('BÀI ÔN TẬP TỰ LUẬN', dam=True, co=15, giua=True)
        v.doan('CHƯƠNG %d: %s' % (so, hoa(goi['ten_chuong'])), dam=True, giua=True)
        v.doan('Học phần: %s (%s) – Học liệu số, giảng dạy trực tuyến' % (TEN_HP, MA_HP), giua=True)
        v.doan('Giảng viên biên soạn: ' + GIANG_VIEN, nghieng=True, giua=True)
        if ban_gv:
            v.doan('BẢN CÓ ĐÁP ÁN GỢI Ý – DÙNG CHO GIẢNG VIÊN', dam=True, giua=True)
            v.doan('HƯỚNG DẪN SỬ DỤNG BẢN NÀY', dam=True, sau=4)
            v.doan('Đây là bản có đáp án gợi ý, để giảng viên đối chiếu khi chấm hoặc thảo luận — không '
                   'phát cho sinh viên (bản phát cho sinh viên là tệp "BAN DE (SV)"). Đáp án là gợi ý: '
                   'bài làm có lập luận khác nhưng đủ các ý ở mục "Tự kiểm" vẫn được chấp nhận.')
        else:
            v.doan('HƯỚNG DẪN LÀM BÀI', dam=True, sau=4)
            v.doan('Đây là bài ôn tập tự luận để tự đánh giá mức độ nắm bài sau khi xem hết hai video '
                   'Chương %d, không phải bài nộp lấy điểm. Mỗi câu nên viết từ nửa trang đến một trang; '
                   'viết xong, dùng danh sách "Tự kiểm những ý cần có" dưới mỗi câu để tự soát bài.' % so)
        muc_truoc = None
        for k, c in enumerate(goi['on_tap_tu_luan'], 1):
            if c['muc'] != muc_truoc:
                v.doan(c['muc'], dam=True, co=14, sau=4)
                muc_truoc = c['muc']
            v.doan('Câu %d. %s' % (k, c['cau_hoi']), dam=True, sau=4)
            v.doan('Tự kiểm những ý cần có:', nghieng=True, sau=2)
            for y in c['tu_kiem']:
                v.doan('– ' + y, thut=12, sau=2)
            if ban_gv:
                v.doan('Đáp án gợi ý:', nghieng=True, sau=2)
                for doan in c['dap_an']:
                    v.doan(doan, sau=4)
            v.doan('', sau=2)
        v.luu()


CHI_VAN_BAN = False


def main():
    global CHI_VAN_BAN
    # --chi-van-ban: chỉ viết lại tệp nội dung slide, kịch bản, ôn tập — bỏ
    # qua bước xuất PDF/ảnh/video (chậm, vài phút mỗi chương).
    CHI_VAN_BAN = '--chi-van-ban' in sys.argv
    so_chuong = [int(x) for x in sys.argv[1:] if x.isdigit()] or [1, 2, 3, 4, 5]
    slides_cd = json.load(open(os.path.join(DOCS, 'slide_noi_dung_chu_de.json'), encoding='utf-8'))
    kich_ban = json.load(open(os.path.join(DOCS, 'kich_ban_chu_de.json'), encoding='utf-8'))
    for so in so_chuong:
        goi = json.load(open(os.path.join(DOCS, 'goi_chuong_ch%02d.json' % so), encoding='utf-8'))
        ban_ghi = dung_chuong(so, slides_cd, kich_ban, goi)
        for video in (1, 2):
            ds = [b for b in ban_ghi if b[2] == video]
            giay = sum(thoi_luong(b[1]) for b in ds)
            ok = 'OK' if 7 * 60 <= giay <= 15 * 60 else 'NGOÀI 7-15 PHÚT'
            print('Chương %d video %d: %2d slide, %4d từ, %.1f phút  %s' % (
                so, video, len(ds), sum(len(b[1].split()) for b in ds), giay / 60, ok))


if __name__ == '__main__':
    main()
